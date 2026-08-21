#!/usr/bin/env python3
"""Append a scored idea to the running spreadsheet log, and rebuild the workbook.

Why a log at all: a single scorecard tells you about one idea, but the value compounds
once there are eight of them side by side. Patterns only show up across rows — the same
criterion binding every time, ceilings that never clear, a network that keeps being one
hop from the buyer. That is a spreadsheet's job, not a chat message's.

Records live as one JSON file per idea under <log-dir>/records/. The workbook is rebuilt
from those records, so it can always be regenerated and the formatting stays consistent.
Columns you fill in by hand (Status, Next action, Test result, Test date, Notes) are read
back out of the existing workbook and carried across every rebuild.

Usage
-----
  log_idea.py add --json '<record>'          # append one idea (or --file record.json)
  log_idea.py add --file records.json        # a list of records also works
  log_idea.py add --file r.json --update     # overwrite an existing id
  log_idea.py build                          # rebuild the workbook from all records
  log_idea.py list                           # what's logged so far

Default log directory is ./idea-log; override with --log-dir.

Record schema — only "id", "idea" and "scores" are required, but a row with empty
reasoning is a row you won't trust in three months, so fill in "why":

  {
    "id": "smid-substack",            "idea": "Paid SMID-cap deep-dive Substack",
    "date": "2026-08-21",             "variant_of": "",
    "buyer": "...", "pain": "...", "v1": "...", "price": "...", "channel": "...",
    "verdict": "one line",
    "scores": {"distribution": 2, ... all ten keys ...},
    "why":    {"distribution": "one line of reasoning", ... same ten keys ...},
    "ceiling_arithmetic": "price x units x retention, spelled out",
    "best_case_yr1": 75000, "realistic_yr1": 42000,
    "capital_required": 2000, "months_to_revenue": 7,
    "carries": "...", "kills": "...",
    "test": "...", "test_pass": "...", "test_fail": "...",
    "reshape": "...", "reshape_score": 76.8,
    "flags": "...", "sources": "..."
  }
"""

import argparse
import datetime
import json
import os
import re
import sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

# Criterion order is canonical everywhere - weights sheet, scorecard columns, JSON keys.
CRITERIA = [
    ("distribution", "Distribution", 15, 1),
    ("demand", "Demand intensity", 15, 1),
    ("fmf", "Founder-market fit", 12, -1),
    ("revenue_structure", "Revenue structure", 8, -1),
    ("defensibility", "Defensibility (3yr)", 8, -1),
    ("leverage", "Leverage type", 10, -1),
    ("capital", "Capital required", 7, -1),
    ("time_to_revenue", "Time-to-first-revenue", 10, -1),
    ("tailwind", "Secular tailwind", 5, -1),
    ("ceiling", "Ceiling (GBP 200k/yr)", 10, 2),
]
KEYS = [c[0] for c in CRITERIA]

# Ascending floors, because MATCH(..., 1) needs ascending order to find the right band.
BANDS = [(0, "Drop"), (40, "Weak"), (55, "Marginal"), (70, "Strong"), (85, "Exceptional")]
CAP_BAND = "Marginal"  # a fired gate can never leave a verdict above this
STATUSES = ["Scored", "Testing", "Validated", "Pursuing", "Shelved", "Dropped"]

ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUBHEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BLUE = Font(name=ARIAL, size=10, color="0000FF")
BLACK = Font(name=ARIAL, size=10)
HEAD_FONT = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=ARIAL, size=13, bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GBP = '£#,##0;(£#,##0);-'

# (sheet, column letter, field) that a human fills in and a rebuild must not clobber.
USER_COLUMNS = [
    ("Summary", "E", "status"),
    ("Summary", "O", "next_action"),
    ("Detail", "N", "test_result"),
    ("Detail", "O", "test_date"),
    ("Detail", "T", "notes"),
]

DATA_ROW = 4  # rows 1-3 are title, spacer, header


# --------------------------------------------------------------------------- records

def records_dir(log_dir):
    return os.path.join(log_dir, "records")


def workbook_path(log_dir):
    return os.path.join(log_dir, "business-ideas-scored.xlsx")


def slugify(text):
    slug = re.sub(r"[^a-z0-9]+", "-", str(text).lower()).strip("-")
    return slug[:60] or "idea"


def validate(record):
    if not isinstance(record, dict):
        raise ValueError("each record must be a JSON object")
    if not record.get("idea"):
        raise ValueError("record needs an 'idea'")
    scores = record.get("scores") or {}
    missing = [k for k in KEYS if k not in scores]
    if missing:
        raise ValueError("%s: missing scores: %s" % (record["idea"], ", ".join(missing)))
    for key in KEYS:
        value = scores[key]
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not 0 <= value <= 5:
            raise ValueError("%s: %s must be a number 0-5, got %r" % (record["idea"], key, value))
    unknown = [k for k in scores if k not in KEYS]
    if unknown:
        raise ValueError("%s: unknown criteria: %s" % (record["idea"], ", ".join(sorted(unknown))))

    record.setdefault("id", slugify(record["idea"]))
    record.setdefault("date", datetime.date.today().isoformat())
    record.setdefault("why", {})
    thin = [k for k in KEYS if not str(record["why"].get(k, "")).strip()]
    return record, thin


def load_records(log_dir):
    directory = records_dir(log_dir)
    if not os.path.isdir(directory):
        return []
    out = []
    for name in sorted(os.listdir(directory)):
        if name.endswith(".json"):
            with open(os.path.join(directory, name)) as fh:
                out.append(json.load(fh))
    out.sort(key=lambda r: (r.get("date", ""), r.get("id", "")))
    return out


def read_user_columns(path):
    """Carry hand-typed cells across a rebuild. Losing them would make the log a
    read-only report rather than something worth living in."""
    kept = {}
    if not os.path.exists(path):
        return kept
    try:
        book = load_workbook(path, data_only=True)
    except Exception as exc:  # a corrupt or open file shouldn't block scoring
        print("note: could not read existing workbook (%s); manual columns not preserved"
              % exc, file=sys.stderr)
        return kept
    for sheet_name, column, field in USER_COLUMNS:
        if sheet_name not in book.sheetnames:
            continue
        sheet = book[sheet_name]
        for row in range(DATA_ROW, sheet.max_row + 1):
            idea_id = sheet.cell(row=row, column=1).value
            value = sheet["%s%d" % (column, row)].value
            if idea_id and value not in (None, ""):
                kept.setdefault(str(idea_id), {})[field] = value
    return kept


# ------------------------------------------------------------------------ formatting

def style_header(sheet, row, headers, widths, wrap_from=0):
    for index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=index, value=title)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BOX
        sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1]
    sheet.row_dimensions[row].height = 34
    sheet.freeze_panes = "%s%d" % (get_column_letter(wrap_from + 1), row + 1)


def title(sheet, text, subtitle=""):
    sheet["A1"] = text
    sheet["A1"].font = TITLE_FONT
    if subtitle:
        sheet["A2"] = subtitle
        sheet["A2"].font = Font(name=ARIAL, size=9, italic=True, color="595959")


# ---------------------------------------------------------------------------- sheets

def build_weights(book):
    sheet = book.create_sheet("Weights")
    title(sheet, "Weights, bands and gates",
          "Everything on the other sheets references these cells. Change a weight or a band "
          "floor here and every scored idea re-ranks.")
    for index, header in enumerate(["Criterion", "Key", "Weight", "Gate: fires at score <=" ], start=1):
        cell = sheet.cell(row=3, column=index, value=header)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
    for offset, (key, label, weight, gate) in enumerate(CRITERIA):
        row = 4 + offset
        sheet.cell(row=row, column=1, value=label).font = BLACK
        sheet.cell(row=row, column=2, value=key).font = Font(name=ARIAL, size=9, color="595959")
        weight_cell = sheet.cell(row=row, column=3, value=weight)
        weight_cell.font = BLUE
        weight_cell.fill = INPUT_FILL
        gate_cell = sheet.cell(row=row, column=4, value=gate)
        gate_cell.font = BLUE
        gate_cell.fill = INPUT_FILL
        for column in range(1, 5):
            sheet.cell(row=row, column=column).border = BOX
    total_row = 4 + len(CRITERIA)
    sheet.cell(row=total_row, column=1, value="Total").font = Font(name=ARIAL, size=10, bold=True)
    total = sheet.cell(row=total_row, column=3, value="=SUM(C4:C%d)" % (total_row - 1))
    total.font = Font(name=ARIAL, size=10, bold=True)
    sheet.cell(row=total_row, column=4,
               value="-1 means the criterion has no gate.").font = Font(
        name=ARIAL, size=9, italic=True, color="595959")

    sheet["A17"] = "Bands"
    sheet["A17"].font = Font(name=ARIAL, size=10, bold=True)
    for index, header in enumerate(["Floor", "Band"], start=1):
        cell = sheet.cell(row=18, column=index, value=header)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.border = BOX
    for offset, (floor, name) in enumerate(BANDS):
        row = 19 + offset
        floor_cell = sheet.cell(row=row, column=1, value=floor)
        floor_cell.font = BLUE
        floor_cell.fill = INPUT_FILL
        floor_cell.border = BOX
        sheet.cell(row=row, column=2, value=name).font = BLACK
        sheet.cell(row=row, column=2).border = BOX

    notes = [
        "",
        "Legend",
        "Blue text on a yellow fill = an input you can change. Black = a formula, leave it alone.",
        "Scores on the Scorecard sheet are inputs too - edit one to see the total move. A rebuild",
        "restores them from the saved record, so make a permanent change by re-scoring the idea.",
        "",
        "Gates: distribution or demand at 1 or below, or ceiling at 2 or below, caps the verdict at",
        "Marginal however high the total. A fatal flaw should not be averaged away. The uncapped",
        "band stays visible on the Scorecard sheet, because 'scores 78, capped' and 'scores 44' are",
        "different problems.",
    ]
    for offset, line in enumerate(notes):
        cell = sheet.cell(row=26 + offset, column=1, value=line)
        cell.font = Font(name=ARIAL, size=9, bold=line == "Legend", color="404040")
    for column, width in zip("ABCD", [26, 18, 10, 22]):
        sheet.column_dimensions[column].width = width
    return sheet


def build_scorecard(book, records):
    sheet = book.create_sheet("Scorecard")
    title(sheet, "Scorecard - every criterion, its score, and why",
          "One row per idea. The 'why' is the point: a score with no sentence behind it is a "
          "number you cannot audit later.")

    headers = ["ID", "Idea"]
    widths = [18, 34]
    for _, label, _, _ in CRITERIA:
        headers += [label, "Wtd", "%s - why" % label]
        widths += [8, 7, 46]
    headers += ["TOTAL", "Band (uncapped)", "Gate fired", "VERDICT", "Binding constraint"]
    widths += [9, 15, 10, 18, 20]
    style_header(sheet, 3, headers, widths, wrap_from=2)

    score_columns, weighted_columns = [], []
    for index in range(len(CRITERIA)):
        score_columns.append(get_column_letter(3 + index * 3))
        weighted_columns.append(get_column_letter(4 + index * 3))
    last = 2 + len(CRITERIA) * 3
    total_col = get_column_letter(last + 1)
    uncapped_col = get_column_letter(last + 2)
    gate_col = get_column_letter(last + 3)
    verdict_col = get_column_letter(last + 4)
    band_first, band_last = 19, 18 + len(BANDS)
    # The gate caps the verdict at Marginal, so it only bites once the total reaches the
    # floor of the band immediately above Marginal. Derived rather than hardcoded, so
    # editing the BANDS table cannot silently point this at the wrong row.
    cap_index = next(i for i, (_, name) in enumerate(BANDS) if name == CAP_BAND)
    cap_floor = "Weights!$A$%d" % (band_first + cap_index + 1)

    for offset, record in enumerate(records):
        row = DATA_ROW + offset
        sheet.cell(row=row, column=1, value=record["id"]).font = BLACK
        sheet.cell(row=row, column=2, value=record["idea"]).font = Font(name=ARIAL, size=10, bold=True)
        for index, key in enumerate(KEYS):
            score = sheet["%s%d" % (score_columns[index], row)]
            score.value = record["scores"][key]
            score.font = BLUE
            score.fill = INPUT_FILL
            score.alignment = Alignment(horizontal="center")
            weighted = sheet["%s%d" % (weighted_columns[index], row)]
            weighted.value = "=ROUND(Weights!$C$%d*%s%d/5,1)" % (
                4 + index, score_columns[index], row)
            weighted.number_format = "0.0"
            weighted.font = BLACK
            why = sheet.cell(row=row, column=5 + index * 3,
                             value=record.get("why", {}).get(key, ""))
            why.font = Font(name=ARIAL, size=9)
            why.alignment = Alignment(wrap_text=True, vertical="top")

        total = sheet["%s%d" % (total_col, row)]
        total.value = "=ROUND(%s,1)" % "+".join(
            "%s%d" % (column, row) for column in weighted_columns)
        total.number_format = "0.0"
        total.font = Font(name=ARIAL, size=10, bold=True)

        uncapped = sheet["%s%d" % (uncapped_col, row)]
        uncapped.value = "=INDEX(Weights!$B$%d:$B$%d,MATCH(%s%d,Weights!$A$%d:$A$%d,1))" % (
            band_first, band_last, total_col, row, band_first, band_last)
        uncapped.font = BLACK

        gate = sheet["%s%d" % (gate_col, row)]
        gate.value = "=IF(OR(%s),\"YES\",\"\")" % ",".join(
            "%s%d<=Weights!$D$%d" % (score_columns[i], row, 4 + i) for i in range(len(CRITERIA)))
        gate.font = Font(name=ARIAL, size=10, bold=True, color="C00000")
        gate.alignment = Alignment(horizontal="center")

        verdict = sheet["%s%d" % (verdict_col, row)]
        verdict.value = '=IF(AND(%s%d="YES",%s%d>=%s),"%s (capped)",%s%d)' % (
            gate_col, row, total_col, row, cap_floor, CAP_BAND, uncapped_col, row)
        verdict.font = Font(name=ARIAL, size=10, bold=True)

        constraint = sheet.cell(row=row, column=last + 5, value=binding_constraint(record))
        constraint.font = Font(name=ARIAL, size=9, italic=True)
        sheet.row_dimensions[row].height = 60

    if records:
        end = DATA_ROW + len(records) - 1
        sheet.conditional_formatting.add(
            "%s%d:%s%d" % (total_col, DATA_ROW, total_col, end),
            ColorScaleRule(start_type="num", start_value=35, start_color="F8696B",
                           mid_type="num", mid_value=62, mid_color="FFEB84",
                           end_type="num", end_value=85, end_color="63BE7B"))
    return sheet


def binding_constraint(record):
    """Lowest score; ties broken by the heavier weight, since that is the one costing
    the most points. Written as a value rather than a formula - the scores it reads sit
    in ten non-adjacent columns, and the formula to walk them would be unreadable."""
    weights = {key: weight for key, _, weight, _ in CRITERIA}
    labels = {key: label for key, label, _, _ in CRITERIA}
    key = min(KEYS, key=lambda k: (record["scores"][k], -weights[k]))
    return "%s (%d/5)" % (labels[key], record["scores"][key])


def build_summary(book, records, kept):
    sheet = book.create_sheet("Summary")
    title(sheet, "Business ideas - scoring log",
          "One row per idea, newest last. Scores and bands are pulled live from the "
          "Scorecard sheet. Status and Next action are yours to fill in.")
    headers = ["ID", "Date scored", "Idea", "Variant of", "Status", "VERDICT", "Total /100",
               "Gate", "Binding constraint", "Best case yr1", "Realistic yr1",
               "Capital needed", "Months to £1", "One-line verdict", "Next action"]
    widths = [18, 12, 34, 16, 12, 17, 10, 8, 20, 14, 14, 14, 11, 52, 30]
    style_header(sheet, 3, headers, widths, wrap_from=3)

    last = 2 + len(CRITERIA) * 3
    columns = {
        "verdict": get_column_letter(last + 4),
        "total": get_column_letter(last + 1),
        "gate": get_column_letter(last + 3),
        "constraint": get_column_letter(last + 5),
    }

    for offset, record in enumerate(records):
        row = DATA_ROW + offset
        manual = kept.get(record["id"], {})
        sheet.cell(row=row, column=1, value=record["id"]).font = BLACK
        sheet.cell(row=row, column=2, value=record.get("date", "")).font = BLACK
        sheet.cell(row=row, column=3, value=record["idea"]).font = Font(
            name=ARIAL, size=10, bold=True)
        sheet.cell(row=row, column=4, value=record.get("variant_of", "")).font = BLACK
        status = sheet.cell(row=row, column=5, value=manual.get("status", "Scored"))
        status.font = BLUE
        status.fill = INPUT_FILL

        for index, field in enumerate(["verdict", "total", "gate", "constraint"]):
            cell = sheet.cell(row=row, column=6 + index)
            cell.value = "=INDEX(Scorecard!${0}:${0},MATCH($A{1},Scorecard!$A:$A,0))".format(
                columns[field], row)
            cell.font = Font(name=ARIAL, size=10, bold=field in ("verdict", "total"))
            cell.alignment = Alignment(horizontal="center")
        sheet.cell(row=row, column=7).number_format = "0.0"

        for index, field in enumerate(["best_case_yr1", "realistic_yr1", "capital_required"]):
            cell = sheet.cell(row=row, column=10 + index, value=record.get(field))
            cell.number_format = GBP
            cell.font = BLACK
        months = sheet.cell(row=row, column=13, value=record.get("months_to_revenue"))
        months.number_format = "0"
        months.font = BLACK

        verdict_text = sheet.cell(row=row, column=14, value=record.get("verdict", ""))
        verdict_text.font = Font(name=ARIAL, size=9)
        verdict_text.alignment = Alignment(wrap_text=True, vertical="top")
        action = sheet.cell(row=row, column=15, value=manual.get("next_action", ""))
        action.font = BLUE
        action.fill = INPUT_FILL
        action.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 46

    if records:
        end = DATA_ROW + len(records) - 1
        validation = DataValidation(type="list", formula1='"%s"' % ",".join(STATUSES),
                                    allow_blank=True, showDropDown=False)
        sheet.add_data_validation(validation)
        validation.add("E%d:E%d" % (DATA_ROW, end))
        sheet.conditional_formatting.add(
            "G%d:G%d" % (DATA_ROW, end),
            ColorScaleRule(start_type="num", start_value=35, start_color="F8696B",
                           mid_type="num", mid_value=62, mid_color="FFEB84",
                           end_type="num", end_value=85, end_color="63BE7B"))
    return sheet


def build_detail(book, records, kept):
    sheet = book.create_sheet("Detail")
    title(sheet, "Detail - the premise, the arithmetic, and the test",
          "The premise row matters most: if the buyer or the price is wrong, argue with "
          "that before arguing with the score.")
    headers = ["ID", "Idea", "Buyer", "Pain", "Crude v1", "Price", "Channel",
               "Ceiling arithmetic", "What carries it", "What kills it",
               "Cheapest disproof test", "Pass looks like", "Fail looks like",
               "Test result", "Test date", "Highest-scoring reshape", "Reshape score",
               "Risks & flags", "Sources", "Notes"]
    widths = [18, 30, 30, 34, 30, 20, 30, 52, 40, 40, 46, 30, 30, 30, 12, 46, 10, 40, 32, 34]
    style_header(sheet, 3, headers, widths, wrap_from=2)

    fields = ["buyer", "pain", "v1", "price", "channel", "ceiling_arithmetic",
              "carries", "kills", "test", "test_pass", "test_fail"]
    for offset, record in enumerate(records):
        row = DATA_ROW + offset
        manual = kept.get(record["id"], {})
        sheet.cell(row=row, column=1, value=record["id"]).font = BLACK
        sheet.cell(row=row, column=2, value=record["idea"]).font = Font(
            name=ARIAL, size=10, bold=True)
        for index, field in enumerate(fields):
            cell = sheet.cell(row=row, column=3 + index, value=record.get(field, ""))
            cell.font = Font(name=ARIAL, size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column, field in ((14, "test_result"), (15, "test_date")):
            cell = sheet.cell(row=row, column=column, value=manual.get(field, ""))
            cell.font = BLUE
            cell.fill = INPUT_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column, field in ((16, "reshape"), (18, "flags"), (19, "sources")):
            cell = sheet.cell(row=row, column=column, value=record.get(field, ""))
            cell.font = Font(name=ARIAL, size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        reshape_score = sheet.cell(row=row, column=17, value=record.get("reshape_score"))
        reshape_score.number_format = "0.0"
        reshape_score.font = BLACK
        notes = sheet.cell(row=row, column=20, value=manual.get("notes", ""))
        notes.font = BLUE
        notes.fill = INPUT_FILL
        notes.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 92
    return sheet


def build(log_dir):
    records = load_records(log_dir)
    path = workbook_path(log_dir)
    kept = read_user_columns(path)

    book = Workbook()
    book.remove(book.active)
    build_summary(book, records, kept)
    build_scorecard(book, records)
    build_detail(book, records, kept)
    build_weights(book)

    os.makedirs(log_dir, exist_ok=True)
    book.save(path)
    return path, records


# ------------------------------------------------------------------------------ cli

def cmd_add(args):
    payload = args.json
    if args.file:
        with open(args.file) as fh:
            payload = fh.read()
    if not payload:
        sys.exit("pass --json or --file")
    try:
        data = json.loads(payload)
    except json.JSONDecodeError as exc:
        sys.exit("input is not valid JSON: %s" % exc)

    incoming = data if isinstance(data, list) else [data]
    os.makedirs(records_dir(args.log_dir), exist_ok=True)
    existing = {r["id"] for r in load_records(args.log_dir)}

    written = []
    for record in incoming:
        try:
            record, thin = validate(record)
        except ValueError as exc:
            sys.exit("bad record: %s" % exc)
        if record["id"] in existing and not args.update:
            sys.exit("'%s' is already logged. Pass --update to overwrite it, or give the "
                     "new version a different id (e.g. '%s-v2')." % (record["id"], record["id"]))
        target = os.path.join(records_dir(args.log_dir), "%s.json" % record["id"])
        with open(target, "w") as fh:
            json.dump(record, fh, indent=2, ensure_ascii=False)
        written.append(record["id"])
        if thin:
            print("note: no reasoning recorded for %s on '%s' - a score with no sentence "
                  "behind it is hard to audit later." % (", ".join(thin), record["id"]),
                  file=sys.stderr)

    path, records = build(args.log_dir)
    print("Logged %s. %d idea%s in %s" % (
        ", ".join(written), len(records), "" if len(records) == 1 else "s", path))


def cmd_build(args):
    path, records = build(args.log_dir)
    print("Rebuilt %s from %d record%s" % (path, len(records), "" if len(records) == 1 else "s"))


def cmd_list(args):
    records = load_records(args.log_dir)
    if not records:
        print("Nothing logged yet in %s" % args.log_dir)
        return
    print("%-20s %-12s %-44s %s" % ("ID", "Date", "Idea", "Binding constraint"))
    print("-" * 104)
    for record in records:
        idea = record["idea"]
        print("%-20s %-12s %-44s %s" % (
            record["id"][:20], record.get("date", ""),
            idea if len(idea) <= 44 else idea[:41] + "...",
            binding_constraint(record)))


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--log-dir", default="idea-log",
                        help="where records and the workbook live (default: ./idea-log)")
    sub = parser.add_subparsers(dest="command", required=True)

    add = sub.add_parser("add", help="append one or more scored ideas")
    source = add.add_mutually_exclusive_group()
    source.add_argument("--json", help="inline JSON record, or a list of them")
    source.add_argument("--file", help="path to a JSON record or list")
    add.add_argument("--update", action="store_true", help="overwrite an existing id")
    add.set_defaults(func=cmd_add)

    sub.add_parser("build", help="rebuild the workbook from saved records").set_defaults(
        func=cmd_build)
    sub.add_parser("list", help="show what has been logged").set_defaults(func=cmd_list)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
